"""Checks for Word documents (.docx).

Every check returns findings already routed. Route 1 is reserved for the two
defects whose repair invents no meaning: a missing document title and a missing
document language. Everything else is a proposal for the person who knows the
content, because marking a row as a header or a paragraph as a heading asserts
something about what the author meant.
"""

from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.document import Document as DocumentType
from docx.oxml.ns import qn
from docx.text.paragraph import Paragraph

from wcag_tool import rules
from wcag_tool.alt import is_missing_or_placeholder
from wcag_tool.findings import Finding, Report

# Both lists come from rules/coverage.json; edit them there.
EVALUATED = rules.load().evaluated("docx")
NOT_EVALUATED = rules.load().not_evaluated("docx")

GENERIC_LINK_TEXT = {"click here", "here", "link", "read more", "more", "this", "this link"}
_HEADING_RE = re.compile(r"^heading (\d)$", re.IGNORECASE)
_MAX_HEADING_WORDS = 12


def style_name(paragraph: Paragraph) -> str:
    """Return the paragraph's style name, or an empty string when it has no style."""
    return paragraph.style.name if paragraph.style is not None else ""


def _heading_level(style_name: str) -> int | None:
    """Return the heading level for a style name such as ``Heading 2``, else None."""
    match = _HEADING_RE.match(style_name.strip())
    return int(match.group(1)) if match else None


def _default_language(doc: DocumentType) -> str | None:
    """Return the document default language code, or None if unset."""
    styles = doc.styles.element
    langs = styles.xpath("./w:docDefaults/w:rPrDefault/w:rPr/w:lang")
    if not langs:
        return None
    return langs[0].get(qn("w:val"))


def _is_fake_heading(paragraph: Paragraph, next_text: str) -> bool:
    """Heuristic: short, bold, unstyled, no terminal punctuation, followed by body text."""
    text = paragraph.text.strip()
    if not text or len(text.split()) > _MAX_HEADING_WORDS or text[-1] in ".!?;:":
        return False
    name = style_name(paragraph)
    if _heading_level(name) is not None or name == "Title":
        return False
    runs = [r for r in paragraph.runs if r.text.strip()]
    if not runs:
        return False
    style_bold = paragraph.style.font.bold if paragraph.style is not None else None
    bold = all(r.bold or (r.font.bold is None and style_bold) for r in runs)
    return bool(bold) and bool(next_text.strip())


def scan(path: Path) -> Report:
    """Scan one .docx file and return a routed report.

    Args:
        path: Path to the document.

    Returns:
        The report with findings, evaluated and not-evaluated criteria.
    """
    doc = Document(str(path))
    asset = path.name
    report = Report.for_format(asset, "docx")
    f = report.findings

    if not (doc.core_properties.title or "").strip():
        f.append(
            Finding.from_check(
                "docx.missing-title",
                asset=asset,
                locator="core properties",
                evidence=(
                    "The document title property is empty; screen readers announce the file name."
                ),
            )
        )

    if _default_language(doc) is None:
        f.append(
            Finding.from_check(
                "docx.missing-language",
                asset=asset,
                locator="document defaults",
                evidence="No default language is set, so a screen reader must guess pronunciation.",
            )
        )

    for doc_pr in doc.element.body.xpath(".//wp:docPr"):
        descr = doc_pr.get("descr")
        if not is_missing_or_placeholder(descr):
            continue
        name = doc_pr.get("name") or f"image {doc_pr.get('id')}"
        evidence = (
            "The image has no alt text."
            if not (descr or "").strip()
            else (
                "The alt text is a filename or an auto-generated phrase: "
                f"{(descr or '').strip()[:60]!r}."
            )
        )
        f.append(
            Finding.from_check(
                "docx.image-missing-alt", asset=asset, locator=name, evidence=evidence
            )
        )

    paragraphs = doc.paragraphs
    levels: list[tuple[int, int]] = []
    for index, paragraph in enumerate(paragraphs):
        level = _heading_level(style_name(paragraph))
        if level is not None and paragraph.text.strip():
            levels.append((index, level))
            continue
        if level is not None:
            f.append(
                Finding.from_check(
                    "docx.empty-heading",
                    asset=asset,
                    locator=f"paragraph {index}",
                    evidence=(
                        f"A Heading {level} paragraph with no text; it is announced as an "
                        "empty heading."
                    ),
                )
            )
            continue
        next_text = paragraphs[index + 1].text if index + 1 < len(paragraphs) else ""
        if _is_fake_heading(paragraph, next_text):
            f.append(
                Finding.from_check(
                    "docx.fake-heading",
                    asset=asset,
                    locator=f"paragraph {index}",
                    evidence=(
                        "Bold text that reads as a heading but is styled as body: "
                        f"{paragraph.text.strip()[:60]!r}"
                    ),
                    fix_id=f"apply-heading:{index}:2",
                )
            )

    body_paragraphs = sum(1 for p in paragraphs if p.text.strip())
    if not levels and body_paragraphs >= 15:
        f.append(
            Finding.from_check(
                "docx.no-headings",
                asset=asset,
                locator="document",
                evidence=(
                    f"{body_paragraphs} paragraphs and no heading styles; the document cannot "
                    "be navigated by structure."
                ),
            )
        )

    previous = 0
    for index, level in levels:
        if previous and level > previous + 1:
            f.append(
                Finding.from_check(
                    "docx.heading-skip",
                    asset=asset,
                    locator=f"paragraph {index}",
                    evidence=f"Heading level jumps from {previous} to {level}.",
                    proposal=f"Use Heading {previous + 1} here, or restructure the sections above.",
                )
            )
        previous = level

    for t_index, table in enumerate(doc.tables):
        if not table.rows:
            continue
        if table.rows[0]._tr.xpath("./w:trPr/w:tblHeader"):
            continue
        f.append(
            Finding.from_check(
                "docx.table-no-header",
                asset=asset,
                locator=f"table {t_index}",
                evidence=(
                    "No row is marked as a header row, so cells are read without their column "
                    "names."
                ),
                fix_id=f"mark-header-row:{t_index}",
            )
        )

    for h_index, link in enumerate(doc.element.body.xpath(".//w:hyperlink")):
        text = "".join(t.text or "" for t in link.xpath(".//w:t")).strip()
        lowered = text.lower()
        if lowered in GENERIC_LINK_TEXT or lowered.startswith(("http://", "https://", "www.")):
            f.append(
                Finding.from_check(
                    "docx.link-text",
                    asset=asset,
                    locator=f"hyperlink {h_index}",
                    evidence=f"Link text is {text[:60]!r}, which does not say where it goes.",
                )
            )

    return report
