"""Fixes for Word documents (.docx).

Each fix is addressed by a ``fix_id`` that a finding carries. Route 1 fixes are
applied by default; any other fix is applied only when it is named explicitly,
because it asserts something about the author's meaning.
"""

from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.document import Document as DocumentType
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

from wcag_tool.docx_checks import _heading_level, style_name

DEFAULT_LANGUAGE = "en-US"


def _title_from_content(doc: DocumentType, path: Path) -> tuple[str, str]:
    """The first Title or Heading 1 paragraph, else a readable file stem, and where it came from."""
    for paragraph in doc.paragraphs:
        name = style_name(paragraph)
        if (name == "Title" or _heading_level(name) == 1) and paragraph.text.strip():
            return paragraph.text.strip(), f"the first {name} paragraph"
    return re.sub(r"[_\-]+", " ", path.stem).strip(), "the file name"


def _set_title(doc: DocumentType, path: Path, title: str | None = None) -> str:
    if title:
        doc.core_properties.title = title
        return f"title set to {title!r}, as named"
    guessed, source = _title_from_content(doc, path)
    doc.core_properties.title = guessed
    return f"title set to {guessed!r} from {source}; check it names the whole document"


def _set_language(doc: DocumentType, code: str = DEFAULT_LANGUAGE) -> str:
    styles = doc.styles.element
    defaults = styles.find(qn("w:docDefaults"))
    if defaults is None:
        defaults = OxmlElement("w:docDefaults")
        styles.insert(0, defaults)
    rpr_default = defaults.find(qn("w:rPrDefault"))
    if rpr_default is None:
        rpr_default = OxmlElement("w:rPrDefault")
        defaults.insert(0, rpr_default)
    rpr = rpr_default.find(qn("w:rPr"))
    if rpr is None:
        rpr = OxmlElement("w:rPr")
        rpr_default.append(rpr)
    lang = rpr.find(qn("w:lang"))
    if lang is None:
        lang = OxmlElement("w:lang")
        rpr.append(lang)
    lang.set(qn("w:val"), code)
    return f"default language set to {code}"


def _mark_header_row(doc: DocumentType, table_index: int) -> str:
    table = doc.tables[table_index]
    tr = table.rows[0]._tr
    tr_pr = tr.get_or_add_trPr()
    if not tr_pr.xpath("./w:tblHeader"):
        tr_pr.append(OxmlElement("w:tblHeader"))
    return f"table {table_index}: first row marked as header"


def _apply_heading(doc: DocumentType, paragraph_index: int, level: int) -> str:
    paragraph = doc.paragraphs[paragraph_index]
    paragraph.style = doc.styles[f"Heading {level}"]
    return f"paragraph {paragraph_index}: Heading {level} applied"


def apply(
    path: Path,
    fix_ids: list[str],
    output: Path,
    language: str = DEFAULT_LANGUAGE,
    title: str | None = None,
) -> dict:
    """Apply the named fixes and save to ``output``.

    Args:
        path: The original document. Never overwritten.
        fix_ids: Fix identifiers such as ``set-title`` or ``mark-header-row:0``.
        output: Where the remediated copy is written.
        language: Language code used by ``set-language``.
        title: The title a person named for ``set-title``; without one, it is taken from
            the document and the note says so.

    Returns:
        A dict with ``applied`` (fix id to description) and ``failed`` (fix id to reason).
    """
    doc = Document(str(path))
    applied: dict[str, str] = {}
    failed: dict[str, str] = {}
    for fix_id in fix_ids:
        name, _, args = fix_id.partition(":")
        try:
            if name == "set-title":
                applied[fix_id] = _set_title(doc, path, title)
            elif name == "set-language":
                applied[fix_id] = _set_language(doc, language)
            elif name == "mark-header-row":
                applied[fix_id] = _mark_header_row(doc, int(args))
            elif name == "apply-heading":
                p_index, level = args.split(":")
                applied[fix_id] = _apply_heading(doc, int(p_index), int(level))
            else:
                failed[fix_id] = "unknown fix"
        except (KeyError, IndexError, ValueError) as exc:
            failed[fix_id] = f"{type(exc).__name__}: {exc}"
    doc.save(str(output))
    return {"applied": applied, "failed": failed}
