"""Generate a realistic, fictional course pack for testing the triage in a chat.

The files in ``demo.py`` are minimal on purpose: each defect sits alone, so the
unit tests and the expected results stay exact. A person testing the BoodleBox
bot needs the opposite: a syllabus, a lecture deck, a handout and two PDFs that
read like a real course, with the defects buried among content that is fine,
including a picture whose alt text is good and a table whose first row is not a
header at all.

Nothing here is a real course, a real person, or a real institution. Names are
placeholders, addresses use ``example.edu``, and the phone number is in the
fictional 555-01xx range.
"""

from __future__ import annotations

import textwrap
from io import BytesIO
from pathlib import Path

from docx import Document
from docx.document import Document as DocumentType
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches as DocxInches
from docx.shared import Pt, RGBColor
from docx.text.paragraph import Paragraph
from PIL import Image, ImageDraw, ImageFont
from pptx import Presentation
from pptx.util import Inches
from pptx.util import Pt as PptxPt
from pypdf import PdfWriter
from pypdf.generic import ContentStream, DictionaryObject, NameObject

COURSE = "ENVS 210: Urban Field Methods"
_RED = RGBColor(0xC0, 0x00, 0x00)
_FIXED_DATE = "D:20261019000000Z"

_SCHEDULE = (
    ("1", "Why walk the city? Course overview", "Syllabus; Field Handbook section 1"),
    ("2", "Observation as method", "Course Reader ch. 1"),
    ("3", "Reading the built environment", "Course Reader ch. 3; Site map exercise"),
    ("4", "Field notes and sketch maps", "Field Handbook section 2; Journal entry 1 due"),
    ("5", "Counting: pedestrians, bikes, benches", "Course Reader ch. 4"),
    ("6", "Site visit 1: Riverside corridor", "Site report 1 due"),
    ("7", "Interviews in public space", "Course Reader ch. 5; consent form review"),
    ("8", "Midterm check-in and journal review", "Journal entries 1-4 due"),
    ("9", "Mapping with open data", "City open data portal tutorial"),
    ("10", "Site visit 2: Market district", "Site report 2 due"),
    ("11", "Access and barriers in public space", "Course Reader ch. 7"),
    ("12", "Analysis: from notes to findings", "Field Handbook section 4"),
    ("13", "Project workshop", "Project draft due"),
    ("14", "Thanksgiving break, no class", ""),
    ("15", "Final presentations", "Final project due"),
)
_DUE = {"4", "6", "8", "10", "13", "15"}


# ----- images ---------------------------------------------------------------------------------


def _font(size: int) -> ImageFont.FreeTypeFont | ImageFont.ImageFont:
    return ImageFont.load_default(size=size)


def _png(image: Image.Image) -> bytes:
    out = BytesIO()
    image.save(out, "PNG", optimize=True)
    return out.getvalue()


def grade_chart_png() -> bytes:
    """A bar chart of the grade breakdown, the kind pasted into a syllabus from a spreadsheet."""
    image = Image.new("RGB", (900, 480), "white")
    draw = ImageDraw.Draw(image)
    draw.text((450, 28), "Grade breakdown", fill="black", font=_font(30), anchor="mm")
    parts = (
        ("Field journal", 30, (31, 119, 180)),
        ("Site reports", 25, (255, 127, 14)),
        ("Final project", 30, (44, 160, 44)),
        ("Participation", 15, (214, 39, 40)),
    )
    base, top, left = 420, 70, 90
    draw.line((left, top, left, base), fill="black", width=2)
    draw.line((left, base, 860, base), fill="black", width=2)
    for tick in range(0, 31, 10):
        y = base - (base - top) * tick / 30
        draw.line((left - 6, y, left, y), fill="black", width=2)
        draw.text((left - 12, y), f"{tick}%", fill="black", font=_font(18), anchor="rm")
    for i, (label, value, colour) in enumerate(parts):
        x = left + 40 + i * 190
        draw.rectangle((x, base - (base - top) * value / 30, x + 120, base), fill=colour)
        draw.text((x + 60, base + 24), label, fill="black", font=_font(20), anchor="mm")
    return _png(image)


def site_map_png() -> bytes:
    """A sketch map of a walking route with numbered stations and a colour-only legend."""
    image = Image.new("RGB", (960, 640), (245, 242, 232))
    draw = ImageDraw.Draw(image)
    for x in range(80, 960, 160):
        draw.rectangle((x, 0, x + 26, 640), fill=(210, 205, 195))
    for y in range(90, 640, 150):
        draw.rectangle((0, y, 960, y + 26), fill=(210, 205, 195))
    draw.rectangle((426, 256, 694, 474), fill=(170, 210, 160))
    draw.text((560, 365), "Riverside Park", fill=(40, 80, 40), font=_font(22), anchor="mm")
    draw.rectangle((0, 560, 960, 640), fill=(150, 190, 225))
    route = [(93, 103), (413, 103), (413, 253), (733, 253), (733, 553)]
    draw.line(route, fill=(200, 30, 30), width=8, joint="curve")
    for n, (x, y) in enumerate(route[:4], start=1):
        draw.ellipse((x - 20, y - 20, x + 20, y + 20), fill="white", outline="black", width=3)
        draw.text((x, y), str(n), fill="black", font=_font(22), anchor="mm")
    draw.rectangle((30, 470, 300, 540), fill="white", outline="black")
    draw.line((45, 492, 95, 492), fill=(200, 30, 30), width=8)
    draw.text((105, 492), "Walking route", fill="black", font=_font(18), anchor="lm")
    draw.rectangle((45, 510, 95, 528), fill=(170, 210, 160))
    draw.text((105, 519), "Public green space", fill="black", font=_font(18), anchor="lm")
    return _png(image)


def counts_chart_png() -> bytes:
    """Last year's pedestrian counts by station, as a slide chart with no data table."""
    image = Image.new("RGB", (900, 520), "white")
    draw = ImageDraw.Draw(image)
    draw.text(
        (450, 30),
        "Pedestrians per 15 minutes, by station",
        font=_font(26),
        fill="black",
        anchor="mm",
    )
    base, left = 450, 80
    draw.line((left, 70, left, base), fill="black", width=2)
    draw.line((left, base, 870, base), fill="black", width=2)
    weekday = (42, 88, 131, 64)
    weekend = (75, 140, 96, 38)
    for i in range(4):
        x = left + 50 + i * 190
        draw.rectangle((x, base - weekday[i] * 2.6, x + 60, base), fill=(31, 119, 180))
        draw.rectangle((x + 66, base - weekend[i] * 2.6, x + 126, base), fill=(255, 127, 14))
        draw.text(
            (x + 63, base + 24), f"Station {i + 1}", fill="black", font=_font(20), anchor="mm"
        )
    return _png(image)


def logo_png() -> bytes:
    """A department wordmark: decorative on a title slide, and usually left with no alt text."""
    image = Image.new("RGB", (360, 120), (24, 64, 52))
    draw = ImageDraw.Draw(image)
    draw.text((180, 60), "Environmental Studies", fill="white", font=_font(26), anchor="mm")
    return _png(image)


def journal_photo_png() -> bytes:
    """An example journal page: sketch boxes and ruled lines, standing in for a photograph."""
    image = Image.new("RGB", (720, 480), (250, 248, 240))
    draw = ImageDraw.Draw(image)
    for y in range(60, 470, 28):
        draw.line((360, y, 700, y), fill=(180, 200, 220), width=2)
    draw.rectangle((30, 40, 330, 300), outline=(60, 60, 60), width=3)
    draw.line((60, 260, 300, 80), fill=(60, 60, 60), width=3)
    draw.text((40, 320), "10:42  bench, shade, 3 people", fill=(40, 40, 40), font=_font(20))
    return _png(image)


# ----- Word -----------------------------------------------------------------------------------


def _strip_language(doc: DocumentType) -> None:
    for lang in doc.styles.element.xpath("./w:docDefaults/w:rPrDefault/w:rPr/w:lang"):
        lang.getparent().remove(lang)


def _link(paragraph: Paragraph, url: str, text: str) -> None:
    """Append a hyperlink styled the way Word styles one: blue and underlined."""
    r_id = paragraph.part.relate_to(url, RT.HYPERLINK, is_external=True)
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), r_id)
    run = OxmlElement("w:r")
    props = OxmlElement("w:rPr")
    colour = OxmlElement("w:color")
    colour.set(qn("w:val"), "0563C1")
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    props.extend([colour, underline])
    t = OxmlElement("w:t")
    t.text = text
    t.set(qn("xml:space"), "preserve")
    run.extend([props, t])
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


def _set_alt(inline_shape, text: str) -> None:  # type: ignore[no-untyped-def]
    inline_shape._inline.docPr.set("descr", text)


def make_syllabus(path: Path) -> Path:
    """A three-page syllabus built the way most syllabi are: partly styled, partly bolded.

    Args:
        path: Where to write the file.

    Returns:
        The path written.
    """
    doc = Document()
    doc.core_properties.title = ""
    doc.core_properties.author = "Course instructor"
    _strip_language(doc)

    title = doc.add_paragraph()
    run = title.add_run(COURSE)
    run.bold = True
    run.font.size = Pt(18)
    doc.add_paragraph("Fall 2026 · Tuesdays and Thursdays, 10:00-11:15 · Hollis Hall 214")

    info = doc.add_table(rows=4, cols=2)
    info.style = "Table Grid"
    for row, (label, value) in enumerate(
        (
            ("Instructor", "Dr. Jordan Example"),
            ("Email", "jexample@example.edu (I reply within one working day)"),
            ("Office hours", "Wednesdays 1:00-3:00, Hollis Hall 309, or by appointment"),
            ("Teaching assistant", "M. Sample, msample@example.edu"),
        )
    ):
        info.cell(row, 0).text = label
        info.cell(row, 1).text = value

    doc.add_heading("Course Description", level=1)
    doc.add_paragraph(
        "This course teaches the methods urban researchers use to study places on foot: "
        "structured observation, field notes and sketch maps, pedestrian counts, short "
        "interviews in public space, and the use of open municipal data. Most weeks include "
        "time outside. You will build a field journal across the term and finish with a "
        "small research project on a site of your choice."
    )

    doc.add_heading("Learning Outcomes", level=1)
    doc.add_paragraph("By the end of the course you will be able to:")
    for outcome in (
        "Design a structured observation protocol for a public space.",
        "Keep field notes that another researcher could follow.",
        "Collect and summarise simple count data.",
        "Conduct a short, consented interview in public and record it ethically.",
        "Present findings about a site in writing and in a short talk.",
    ):
        doc.add_paragraph(outcome, style="List Number")

    doc.add_heading("Required Materials", level=3)
    doc.add_paragraph(
        "The Course Reader (available in the bookstore and as a PDF on the course page), the "
        "Field Handbook (free on the course page), a bound notebook you can write in outdoors, "
        "and shoes you can walk two miles in."
    )

    policies = doc.add_paragraph()
    policies.add_run("Course Policies").bold = True
    doc.add_paragraph(
        "Attendance. We meet outside on most Thursdays, and fieldwork cannot be made up from a "
        "recording. If you will miss a field session, email me before it so we can arrange an "
        "alternative site."
    )
    doc.add_paragraph(
        "Late work. Site reports lose one letter grade for each day late unless an extension "
        "is agreed in advance. Journal entries are accepted up to one week late without "
        "penalty."
    )
    doc.add_paragraph(
        "Weather. If the campus closes, the field session moves online. Check the course page "
        "by 8:00 on the day."
    )
    doc.add_paragraph("Items marked in red on the schedule below are due before class that week.")

    doc.add_heading("Weekly Schedule", level=1)
    schedule = doc.add_table(rows=len(_SCHEDULE) + 1, cols=3)
    schedule.style = "Table Grid"
    for col, text in enumerate(("Week", "Topic", "Reading and work due")):
        cell = schedule.cell(0, col)
        cell.text = ""
        cell.paragraphs[0].add_run(text).bold = True
    for row, (week, topic, due) in enumerate(_SCHEDULE, start=1):
        schedule.cell(row, 0).text = week
        schedule.cell(row, 1).text = topic
        due_run = schedule.cell(row, 2).paragraphs[0].add_run(due)
        if week in _DUE:
            due_run.font.color.rgb = _RED
    doc.add_heading("", level=2)

    doc.add_heading("Grading", level=1)
    grading = doc.add_table(rows=5, cols=2)
    grading.style = "Table Grid"
    for row, (part, weight) in enumerate(
        (
            ("Component", "Weight"),
            ("Field journal (weekly entries)", "30%"),
            ("Site reports (two)", "25%"),
            ("Final project and presentation", "30%"),
            ("Participation in field sessions", "15%"),
        )
    ):
        grading.cell(row, 0).text = part
        grading.cell(row, 1).text = weight
    doc.add_paragraph()
    doc.add_picture(BytesIO(grade_chart_png()), width=DocxInches(5.5))
    doc.add_paragraph("Figure 1. Grade breakdown.", style="Caption")

    doc.add_heading("Accessibility and Support", level=1)
    support = doc.add_paragraph(
        "If you have a disability and need an accommodation in this course, contact the "
        "Accessibility Resource Office as early in the term as you can. To request one, "
    )
    _link(support, "https://example.edu/accessibility/request", "click here")
    support.add_run(
        ". All field sites have been chosen with step-free routes; tell me if a "
        "route does not work for you and we will find another."
    )
    library = doc.add_paragraph("Library guide for this course: ")
    _link(
        library,
        "https://example.edu/library/guides/envs210",
        "https://example.edu/library/guides/envs210",
    )
    writing = doc.add_paragraph("The Writing Center reviews drafts of site reports. Book a time ")
    _link(writing, "https://example.edu/writing-center/book", "here")
    writing.add_run(".")
    doc.add_paragraph(
        "Questions about the course can go to my email or to the course discussion board. "
        "For an emergency on a field day, call the department office at 555-0142."
    )

    doc.save(str(path))
    return path


def make_journal_guidelines(path: Path) -> Path:
    """A long handout with its title and language set, a good image, and no headings at all.

    Args:
        path: Where to write the file.

    Returns:
        The path written.
    """
    doc = Document()
    doc.core_properties.title = "Field Journal Guidelines"
    doc.core_properties.author = "Course instructor"
    paragraphs = (
        "FIELD JOURNAL GUIDELINES",
        "ENVS 210, Fall 2026",
        "Your field journal is the record of your fieldwork this term. It is graded on how "
        "complete and honest it is, not on how polished it looks. A messy page with good "
        "notes is worth more than a tidy page written up afterwards from memory.",
        "WHAT TO BRING",
        "A bound notebook, two pencils, a watch or phone for timing, and a clipboard if you "
        "have one. Pens smudge in rain. Loose sheets get lost.",
        "BEFORE YOU START",
        "Write the date, the start time, the weather, and the name of the site at the top of "
        "the page. Sketch the site from where you are standing before you write anything else.",
        "Mark where you are standing on the sketch, and which way you are facing.",
        "DURING OBSERVATION",
        "Write what you see, not what you think it means. 'Three people sitting on the low "
        "wall in the shade' is an observation. 'People prefer the wall' is an interpretation, "
        "and it belongs in the right-hand column.",
        "Time-stamp every entry. If nothing happens for ten minutes, write that down too.",
        "Count, do not estimate. If there are too many people to count, count one section of "
        "the space and say which section.",
        "Note who is not there. A plaza with no older adults or no children at 4 p.m. on a "
        "weekday is a finding.",
        "INTERVIEWS",
        "Only interview people who have agreed after you have read them the consent script. "
        "Do not write down names, and do not photograph faces.",
        "AFTER EACH SESSION",
        "Within 24 hours, write a half-page reflection: what surprised you, what you would "
        "do differently, and one question the session raised.",
        "Transfer your counts to the shared spreadsheet on the course page so the class can "
        "compare sites.",
        "HOW IT IS GRADED",
        "Each entry is marked complete, partial or missing. Complete entries have a sketch, "
        "time stamps, at least ten observations, and a reflection. Partial entries are "
        "missing one of these.",
        "You may submit entries up to one week late without penalty.",
        "An example of a complete entry is shown below.",
    )
    for text in paragraphs:
        doc.add_paragraph(text)
    picture = doc.add_picture(BytesIO(journal_photo_png()), width=DocxInches(4.5))
    _set_alt(
        picture,
        "A sample journal page: a sketch map of a small plaza on the left, ruled notes with "
        "time stamps on the right.",
    )
    doc.add_paragraph("Questions? Bring them to office hours or post them on the course board.")
    doc.save(str(path))
    return path


# ----- PowerPoint -----------------------------------------------------------------------------


def _picture_alt(picture, text: str) -> None:  # type: ignore[no-untyped-def]
    picture._element.xpath(".//p:cNvPr")[0].set("descr", text)


def _bullets(slide, items: list[tuple[int, str]]) -> None:  # type: ignore[no-untyped-def]
    frame = slide.placeholders[1].text_frame
    frame.text = items[0][1]
    for level, text in items[1:]:
        paragraph = frame.add_paragraph()
        paragraph.text = text
        paragraph.level = level


def _notes(slide, text: str) -> None:  # type: ignore[no-untyped-def]
    slide.notes_slide.notes_text_frame.text = text


def make_lecture_deck(path: Path) -> Path:
    """A seven-slide lecture deck with the defects decks really have.

    A logo with no alt text on the title slide, a map whose alt text is the file
    name, a chart with an auto-generated description, a slide titled with a text
    box instead of the title placeholder, and a table with no header row.

    Args:
        path: Where to write the file.

    Returns:
        The path written.
    """
    prs = Presentation()
    prs.core_properties.title = ""
    prs.core_properties.author = "Course instructor"
    layouts = {layout.name: layout for layout in prs.slide_layouts}

    slide = prs.slides.add_slide(layouts["Title Slide"])
    slide.shapes.title.text = "Week 3: Reading the Built Environment"
    slide.placeholders[1].text = "ENVS 210: Urban Field Methods · Fall 2026"
    logo = slide.shapes.add_picture(BytesIO(logo_png()), Inches(0.4), Inches(6.4), Inches(2.4))
    _picture_alt(logo, "")
    _notes(slide, "Welcome back. Collect journal entry 1 sign-up sheet at the door.")

    slide = prs.slides.add_slide(layouts["Title and Content"])
    slide.shapes.title.text = "Today"
    _bullets(
        slide,
        [
            (0, "Paths, edges, districts, nodes and landmarks"),
            (0, "How to see them on a walk"),
            (0, "Thursday's route and what to record"),
            (0, "Last year's counts: what they tell us"),
        ],
    )

    slide = prs.slides.add_slide(layouts["Title and Content"])
    slide.shapes.title.text = "Five elements of a city's image"
    _bullets(
        slide,
        [
            (0, "Paths: the channels people move along"),
            (0, "Edges: boundaries that are not paths, such as a river or a rail line"),
            (0, "Districts: areas with a recognisable character"),
            (0, "Nodes: points you enter, such as junctions and squares"),
            (0, "Landmarks: reference points seen from outside"),
            (1, "Course Reader ch. 3 develops each with examples"),
        ],
    )
    _notes(slide, "Ask: which of these can you see from the classroom window?")

    slide = prs.slides.add_slide(layouts["Title Only"])
    slide.shapes.title.text = "Thursday's route"
    route = slide.shapes.add_picture(BytesIO(site_map_png()), Inches(1.0), Inches(1.6), Inches(8.0))
    _picture_alt(route, "IMG_2041.png")
    _notes(slide, "Meet at station 1. The red line is the route; stop at each numbered circle.")

    slide = prs.slides.add_slide(layouts["Blank"])
    box = slide.shapes.add_textbox(Inches(0.5), Inches(0.4), Inches(9.0), Inches(1.0))
    box.text_frame.text = "What to record at each station"
    heading = box.text_frame.paragraphs[0].runs[0].font
    heading.size = PptxPt(36)
    heading.bold = True
    table = slide.shapes.add_table(5, 3, Inches(0.5), Inches(1.6), Inches(9.0), Inches(3.6)).table
    table.first_row = False
    for row, cells in enumerate(
        (
            ("Station", "Record", "Time"),
            ("1. Bridge Street junction", "Pedestrian count, crossing behaviour", "15 min"),
            ("2. Riverside Park gate", "Who enters, and how", "15 min"),
            ("3. Market corner", "Edges and seating, sketch map", "20 min"),
            ("4. Ferry steps", "Landmarks visible, interview if consented", "20 min"),
        )
    ):
        for col, text in enumerate(cells):
            table.cell(row, col).text = text

    slide = prs.slides.add_slide(layouts["Title Only"])
    slide.shapes.title.text = "Last year's counts"
    chart = slide.shapes.add_picture(
        BytesIO(counts_chart_png()), Inches(0.8), Inches(1.6), Inches(8.4)
    )
    _picture_alt(chart, "A picture containing chart Description automatically generated")
    _notes(slide, "Blue is weekday, orange is weekend. Station 3 flips. Why?")

    slide = prs.slides.add_slide(layouts["Title and Content"])
    slide.shapes.title.text = "Before Thursday"
    _bullets(
        slide,
        [
            (0, "Read Course Reader ch. 3"),
            (0, "Print the consent script, or have it on your phone"),
            (0, "Check the weather; we go unless campus closes"),
            (0, "Questions: see the discussion board link on the course page"),
        ],
    )

    prs.save(str(path))
    return path


# ----- PDF ------------------------------------------------------------------------------------


def _pdf_escape(text: str) -> bytes:
    escaped = text.replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)")
    return escaped.encode("cp1252")


def _text_page(writer: PdfWriter, lines: list[tuple[str, int, bool]]) -> None:
    """Add a letter page with lines of (text, size, bold), top to bottom."""
    page = writer.add_blank_page(width=612, height=792)
    fonts = DictionaryObject()
    for key, base in (("/F1", "/Helvetica"), ("/F2", "/Helvetica-Bold")):
        fonts[NameObject(key)] = DictionaryObject(
            {
                NameObject("/Type"): NameObject("/Font"),
                NameObject("/Subtype"): NameObject("/Type1"),
                NameObject("/BaseFont"): NameObject(base),
                NameObject("/Encoding"): NameObject("/WinAnsiEncoding"),
            }
        )
    page[NameObject("/Resources")] = DictionaryObject({NameObject("/Font"): fonts})
    y = 732
    data = b""
    for text, size, bold in lines:
        if text:
            font = b"/F2" if bold else b"/F1"
            data += b"BT %s %d Tf 72 %d Td (%s) Tj ET\n" % (font, size, y, _pdf_escape(text))
        y -= int(size * 1.45) if text else 10
    content = ContentStream(None, writer)
    content.set_data(data)
    page.replace_contents(content)


def _wrapped(text: str, width: int = 88) -> list[tuple[str, int, bool]]:
    return [(line, 11, False) for line in textwrap.wrap(text, width)]


def make_reading_list_pdf(path: Path) -> Path:
    """A two-page reading list exported from Word without tags, title or language.

    Args:
        path: Where to write the file.

    Returns:
        The path written.
    """
    writer = PdfWriter()
    first: list[tuple[str, int, bool]] = [
        ("ENVS 210: Urban Field Methods", 16, True),
        ("Week 3 Reading List and Questions", 13, True),
        ("", 0, False),
        ("Required", 12, True),
        *_wrapped(
            "1. Course Reader, chapter 3: 'Paths, Edges and Districts' (pp. 41-66). Read the "
            "whole chapter. Skim the case on pp. 58-61 if you are short of time."
        ),
        *_wrapped(
            "2. Field Handbook, section 2: 'Field Notes and Sketch Maps' (pp. 9-17). Bring the "
            "handbook to Thursday's session."
        ),
        ("", 0, False),
        ("Recommended", 12, True),
        *_wrapped(
            "3. City Planning Department (2024). Riverside Corridor Pedestrian Study, summary "
            "pages 1-6. On the course page under Week 3."
        ),
        *_wrapped(
            "4. Course Reader, chapter 4: 'Counting People in Public Space' (pp. 67-80). We "
            "start this in Week 5; reading it now will help with Thursday's counts."
        ),
        ("", 0, False),
        ("Questions to bring to Tuesday's class", 12, True),
        *_wrapped(
            "a. Which of the five elements is hardest to see on the ground rather than on a map? "
            "Give one example from somewhere you walk every week."
        ),
        *_wrapped(
            "b. Chapter 3 argues that edges can join areas as well as divide them. Find one edge "
            "on campus that does each."
        ),
        *_wrapped(
            "c. Look at the figure on p. 52 (the shaded district map). What would a person who "
            "cannot see the shading need to be told to understand it?"
        ),
    ]
    second: list[tuple[str, int, bool]] = [
        ("Reading tips", 12, True),
        *_wrapped(
            "Read with a pencil. Mark the two places where the author makes a claim you could "
            "test on Thursday's walk, and write the test in the margin."
        ),
        *_wrapped(
            "The PDF of the Course Reader is a scan of the printed edition. If it does not work "
            "with your screen reader or text-to-speech tool, email me and I will send the text."
        ),
        ("", 0, False),
        ("Where to find the readings", 12, True),
        *_wrapped(
            "Course Reader: bookstore, or the PDF on the course page under Readings. Field "
            "Handbook: course page under Handbooks. Pedestrian study: course page under Week 3."
        ),
        ("", 0, False),
        ("ENVS 210 · Fall 2026 · Page 2", 9, False),
    ]
    _text_page(writer, first)
    _text_page(writer, second)
    with path.open("wb") as handle:
        writer.write(handle)
    return path


_SCAN_TEXT = (
    (
        "CHAPTER 3",
        "Paths, Edges and Districts",
    ),
    (
        "A city is learned on foot before it is learned from a map. The first things a "
        "newcomer fixes are the routes they repeat: the walk from a bus stop to a front door, "
        "the short cut across a car park, the street that always has somewhere open to buy "
        "coffee. These paths are the skeleton of a personal city, and they are remarkably "
        "stable once laid down.",
        "Edges are less obvious. A river, a rail cutting or a six-lane road is plainly a "
        "boundary, but so is the point where shopfronts give way to blank walls, or where "
        "the paving changes. People rarely name these lines, yet they route around them, and "
        "they describe the areas on either side as different places.",
        "Districts follow. An area becomes a district when it is recognisable from inside: "
        "the same building heights, the same kind of shop, the same people at the same hour. "
        "The field exercises in this course ask you to test that recognition rather than "
        "assume it.",
    ),
)


def make_scanned_chapter_pdf(path: Path) -> Path:
    """Two scanned pages of a course reader: readable to the eye, silent to a screen reader.

    Args:
        path: Where to write the file.

    Returns:
        The path written.
    """
    pages = []
    for number, (head, body) in enumerate(
        ((_SCAN_TEXT[0], _SCAN_TEXT[1][:2]), (None, _SCAN_TEXT[1][2:])), start=41
    ):
        page = Image.new("L", (1100, 1420), 238)
        draw = ImageDraw.Draw(page)
        y = 120
        if head:
            draw.text((550, y), head[0], fill=40, font=_font(26), anchor="mm")
            draw.text((550, y + 56), head[1], fill=25, font=_font(40), anchor="mm")
            y += 150
        for paragraph in body:
            for line in textwrap.wrap(paragraph, 62):
                draw.text((110, y), line, fill=30, font=_font(28))
                y += 42
            y += 28
        draw.text((550, 1350), str(number), fill=60, font=_font(22), anchor="mm")
        pages.append(page.rotate(0.6, resample=Image.Resampling.BICUBIC, fillcolor=238))
    pages[0].save(
        path,
        "PDF",
        resolution=130.0,
        save_all=True,
        append_images=pages[1:],
        creationDate=_FIXED_DATE,
        modDate=_FIXED_DATE,
    )
    return path


# ----- the pack -------------------------------------------------------------------------------


def make_course_pack(directory: Path) -> list[Path]:
    """Write five fictional course files that read like a real course.

    Args:
        directory: Created if needed.

    Returns:
        The paths written, in a stable order.
    """
    directory.mkdir(parents=True, exist_ok=True)
    return [
        make_syllabus(directory / "ENVS210_Syllabus_Fall2026.docx"),
        make_lecture_deck(directory / "ENVS210_Week03_Lecture.pptx"),
        make_journal_guidelines(directory / "Field_Journal_Guidelines.docx"),
        make_reading_list_pdf(directory / "Week03_Reading_List.pdf"),
        make_scanned_chapter_pdf(directory / "Course_Reader_Ch3_scan.pdf"),
    ]
