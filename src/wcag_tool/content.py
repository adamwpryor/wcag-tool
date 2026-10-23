"""A document's words and objects, in reading order, for an agent to read every page.

The scan reads settings; it cannot read for meaning. An agent can, but only if it has
the words, and an agent kept from running its own code has no other way to get them.
``extract`` returns them: per page (PDF), per slide (PowerPoint), or for the whole
document (Word), as ordered blocks, each with the same place names the scan uses, so
"Paragraph 14" in the work order and here are the same paragraph.

It reports what the file holds, not what it means: a PDF with no tags has no headings
to report, only text, and nothing here guesses a role the file does not mark. It reads
the file and writes nothing. The same blocks are what an agent rebuilds a PDF from;
``compare`` checks the rebuild against them.
"""

from __future__ import annotations

import re
from pathlib import Path
from typing import Any

from docx import Document
from docx.document import Document as DocumentType
from docx.oxml.ns import qn
from docx.table import Table
from docx.text.paragraph import Paragraph
from lxml.etree import _Element
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.shapes.base import BaseShape
from pptx.shapes.group import GroupShape
from pypdf import PdfReader

from wcag_tool.docx_checks import style_name
from wcag_tool.pdf_checks import drawn_images

_HEADING = re.compile(r"^heading (\d)$", re.IGNORECASE)


def _decorative(element: _Element) -> bool:
    """True when Office's *Mark as decorative* is set on a picture's properties."""
    return any(
        str(node.tag).endswith("}decorative") and node.get("val") in ("1", "true")
        for node in element.iter()
    )


# ----- Word -----------------------------------------------------------------------------------


def _docx_links(element: _Element, doc: DocumentType, index: dict[_Element, int]) -> list[dict]:
    links = []
    for link in element.xpath(".//w:hyperlink"):
        text = "".join(t.text or "" for t in link.xpath(".//w:t")).strip()
        rid = link.get(qn("r:id"))
        anchor = link.get(qn("w:anchor"))
        target = doc.part.rels[rid].target_ref if rid in doc.part.rels else f"#{anchor or ''}"
        links.append({"where": f"Link {index[link] + 1}", "text": text, "target": target})
    return links


def _docx_pictures(element: _Element) -> list[dict]:
    pictures = []
    for doc_pr in element.xpath(".//wp:docPr"):
        pictures.append(
            {
                "kind": "picture",
                "where": doc_pr.get("name") or f"image {doc_pr.get('id')}",
                "alt": (doc_pr.get("descr") or "").strip() or None,
                "decorative": _decorative(doc_pr),
            }
        )
    return pictures


def _docx_paragraph(paragraph: Paragraph, number: int) -> dict[str, Any] | None:
    text = paragraph.text.strip()
    boxes = " ".join(
        "".join(t.text or "" for t in box.xpath(".//w:t")).strip()
        for box in paragraph._p.xpath(".//w:txbxContent")
    ).strip()
    name = style_name(paragraph)
    heading = _HEADING.match(name.strip())
    if not text and not boxes:
        # An empty paragraph is spacing; an empty heading is a defect a reader meets.
        if heading:
            level = int(heading.group(1))
            where = f"Paragraph {number}"
            return {"kind": "heading", "where": where, "text": "", "level": level, "empty": True}
        return None
    p_pr = paragraph._p.pPr
    numbered = p_pr is not None and p_pr.numPr is not None
    if heading:
        kind = "heading"
    elif name.lower().startswith("list") or numbered:
        kind = "list_item"
    else:
        kind = "paragraph"
    runs = [r for r in paragraph.runs if r.text.strip()]
    block: dict[str, Any] = {"kind": kind, "where": f"Paragraph {number}", "text": text}
    if heading:
        block["level"] = int(heading.group(1))
    elif name:
        block["style"] = name
    if runs and all(r.bold for r in runs):
        block["bold"] = True
    if boxes:
        block["text_boxes"] = boxes
    return block


def _docx(path: Path) -> dict[str, Any]:
    doc = Document(str(path))
    body = doc.element.body
    # Keyed by element, which keeps lxml's wrappers alive so later queries return the same ones.
    link_index = {link: n for n, link in enumerate(body.xpath(".//w:hyperlink"))}
    blocks: list[dict[str, Any]] = []
    p_number = t_number = 0
    for child in body.iterchildren():
        if child.tag == qn("w:p"):
            p_number += 1
            paragraph = Paragraph(child, doc)
            block = _docx_paragraph(paragraph, p_number)
            if block:
                links = _docx_links(child, doc, link_index)
                if links:
                    block["links"] = links
                blocks.append(block)
            blocks += _docx_pictures(child)
        elif child.tag == qn("w:tbl"):
            t_number += 1
            table = Table(child, doc)
            header = bool(table.rows) and bool(table.rows[0]._tr.xpath("./w:trPr/w:tblHeader"))
            block = {
                "kind": "table",
                "where": f"Table {t_number}",
                "header_row_marked": header,
                "rows": [[cell.text.strip() for cell in row.cells] for row in table.rows],
            }
            links = _docx_links(child, doc, link_index)
            if links:
                block["links"] = links
            blocks.append(block)
            blocks += _docx_pictures(child)
    margins: list[dict[str, Any]] = []
    for section in doc.sections:
        for part, where in ((section.header, "Header"), (section.footer, "Footer")):
            text = " ".join(p.text.strip() for p in part.paragraphs if p.text.strip())
            if text and not any(m["text"] == text for m in margins):
                margins.append({"kind": "margin", "where": where, "text": text})
    langs = doc.styles.element.xpath("./w:docDefaults/w:rPrDefault/w:rPr/w:lang")
    return {
        "settings": {
            "title": doc.core_properties.title or None,
            "language": langs[0].get(qn("w:val")) if langs else None,
        },
        "units": [{"unit": "document", "blocks": blocks + margins}],
    }


# ----- PowerPoint -----------------------------------------------------------------------------


def _shapes(shapes: list[BaseShape]) -> list[BaseShape]:
    out: list[BaseShape] = []
    for shape in shapes:
        if isinstance(shape, GroupShape):
            out += _shapes(list(shape.shapes))
        else:
            out.append(shape)
    return out


def _slide_blocks(slide: Any, number: int) -> list[dict[str, Any]]:
    title = slide.shapes.title
    # python-pptx's shape classes share no typed base for text, tables and pictures.
    shapes: list[Any] = sorted(_shapes(list(slide.shapes)), key=lambda s: (s.top or 0, s.left or 0))
    blocks: list[dict[str, Any]] = []
    for shape in shapes:
        where = f"slide {number}, {shape.name}"
        if title is not None and shape.shape_id == title.shape_id:
            if shape.text.strip():
                blocks.append({"kind": "heading", "level": 1, "where": where, "text": shape.text})
            continue
        if shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
            c_nv_pr = shape._element.xpath(".//p:cNvPr")
            alt = c_nv_pr[0].get("descr") if c_nv_pr else None
            decorative = _decorative(c_nv_pr[0]) if c_nv_pr else False
            blocks.append(
                {
                    "kind": "picture",
                    "where": where,
                    "alt": (alt or "").strip() or None,
                    "decorative": decorative,
                }
            )
        elif shape.has_table:
            table = shape.table
            blocks.append(
                {
                    "kind": "table",
                    "where": where,
                    "header_row_marked": bool(table.first_row),
                    "rows": [[cell.text.strip() for cell in row.cells] for row in table.rows],
                }
            )
        elif shape.has_text_frame:
            for paragraph in shape.text_frame.paragraphs:
                text = "".join(run.text for run in paragraph.runs).strip()
                if not text:
                    continue
                block: dict[str, Any] = {
                    "kind": "list_item" if paragraph.level > 0 else "paragraph",
                    "where": where,
                    "text": text,
                }
                if paragraph.level:
                    block["level"] = paragraph.level
                links = [
                    {"text": run.text.strip(), "target": run.hyperlink.address}
                    for run in paragraph.runs
                    if run.hyperlink.address
                ]
                if links:
                    block["links"] = links
                blocks.append(block)
    if slide.has_notes_slide:
        notes = slide.notes_slide.notes_text_frame
        if notes is not None and notes.text.strip():
            blocks.append({"kind": "notes", "where": f"slide {number}, notes", "text": notes.text})
    return blocks


def _pptx(path: Path) -> dict[str, Any]:
    prs = Presentation(str(path))
    return {
        "settings": {"title": prs.core_properties.title or None},
        "units": [
            {"unit": f"slide {n}", "blocks": _slide_blocks(slide, n)}
            for n, slide in enumerate(prs.slides, start=1)
        ],
    }


# ----- PDF ------------------------------------------------------------------------------------


def _pdf(path: Path) -> dict[str, Any]:
    reader = PdfReader(str(path))
    root = reader.root_object
    mark = root.get("/MarkInfo")
    marked = mark.get_object().get("/Marked") if mark is not None else None
    title = reader.metadata.title if reader.metadata is not None else None
    lang = root.get("/Lang")
    units = []
    for number, page in enumerate(reader.pages, start=1):
        blocks: list[dict[str, Any]] = []
        text = (page.extract_text() or "").strip()
        if text:
            blocks.append({"kind": "text", "where": f"page {number}", "text": text})
        pictures = len(drawn_images(reader, page))
        if pictures:
            blocks.append({"kind": "pictures", "where": f"page {number}", "count": pictures})
        for ref in page.get("/Annots") or []:
            annotation = ref.get_object()
            subtype = annotation.get("/Subtype")
            if subtype == "/Link" and "/A" in annotation:
                uri = annotation["/A"].get_object().get("/URI")
                if uri:
                    blocks.append({"kind": "link", "where": f"page {number}", "target": str(uri)})
            elif subtype == "/Widget":
                field, name, tooltip, kind = annotation, None, None, None
                while field is not None:
                    field = field.get_object()
                    name = name or field.get("/T")
                    tooltip = tooltip or field.get("/TU")
                    kind = kind or field.get("/FT")
                    field = field.get("/Parent")
                blocks.append(
                    {
                        "kind": "field",
                        "where": f"page {number}",
                        "name": str(name) if name else None,
                        "tooltip": str(tooltip) if tooltip else None,
                        "type": {"/Tx": "text", "/Btn": "checkbox or button"}.get(
                            str(kind), str(kind)
                        ),
                    }
                )
        units.append({"unit": f"page {number}", "blocks": blocks})
    return {
        "settings": {
            "title": title or None,
            "language": str(lang) if lang else None,
            "tagged": bool(marked),
        },
        "units": units,
    }


def extract(path: Path) -> dict[str, Any]:
    """A document's words and objects, in reading order, with the scan's place names.

    Args:
        path: The course file. Read, never written.

    Returns:
        ``file``, ``format``, ``settings`` (title, language, and for a PDF whether it is
        tagged), and ``units``: one per page, slide, or the whole Word document, each
        with ordered ``blocks`` (heading, paragraph, list item, table, picture, link,
        field, notes, header or footer text).

    Raises:
        ValueError: For a file type other than .docx, .pptx or .pdf.
    """
    suffix = path.suffix.lower()
    readers = {".docx": _docx, ".pptx": _pptx, ".pdf": _pdf}
    if suffix not in readers:
        raise ValueError(f"Unsupported file type {suffix!r}: expected .docx, .pptx or .pdf.")
    return {"file": path.name, "format": suffix[1:], **readers[suffix](path)}
