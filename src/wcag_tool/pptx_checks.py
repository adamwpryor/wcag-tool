"""Checks and fixes for PowerPoint decks (.pptx).

Reading order is the defect that matters most in slides and the one the tool
cannot judge, so it is listed under not-evaluated with instructions rather than
guessed at.
"""

from __future__ import annotations

from pathlib import Path

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

from wcag_tool import rules
from wcag_tool.alt import is_missing_or_placeholder
from wcag_tool.findings import Finding, Report

# Both lists come from rules/coverage.json; edit them there.
EVALUATED = rules.load().evaluated("pptx")
NOT_EVALUATED = rules.load().not_evaluated("pptx")


def scan(path: Path) -> Report:
    """Scan one .pptx file and return a routed report.

    Args:
        path: Path to the deck.

    Returns:
        The report.
    """
    prs = Presentation(str(path))
    asset = path.name
    report = Report.for_format(asset, "pptx")
    f = report.findings

    if not (prs.core_properties.title or "").strip():
        f.append(
            Finding.from_check(
                "pptx.missing-title",
                asset=asset,
                locator="core properties",
                evidence="The presentation title property is empty.",
            )
        )

    for number, slide in enumerate(prs.slides, start=1):
        title_shape = slide.shapes.title
        if title_shape is None or not title_shape.has_text_frame or not title_shape.text.strip():
            f.append(
                Finding.from_check(
                    "pptx.slide-no-title",
                    asset=asset,
                    locator=f"slide {number}",
                    evidence=(
                        "No title placeholder with text; screen readers announce the slide by "
                        "number only."
                    ),
                )
            )
        for shape_index, shape in enumerate(slide.shapes):
            if shape.has_table and not shape.table.first_row:
                f.append(
                    Finding.from_check(
                        "pptx.table-no-header",
                        asset=asset,
                        locator=f"slide {number}, {shape.name}",
                        evidence="The table has no header row marked.",
                        fix_id=f"mark-header-row:{number}:{shape_index}",
                    )
                )
                continue
            if shape.shape_type != MSO_SHAPE_TYPE.PICTURE:
                continue
            c_nv_pr = shape._element.xpath(".//p:cNvPr")
            descr = c_nv_pr[0].get("descr") if c_nv_pr else None
            if not is_missing_or_placeholder(descr):
                continue
            evidence = (
                "The picture has no alt text."
                if not (descr or "").strip()
                else (
                    "The alt text is a filename or an auto-generated phrase: "
                    f"{(descr or '').strip()[:60]!r}."
                )
            )
            f.append(
                Finding.from_check(
                    "pptx.image-missing-alt",
                    asset=asset,
                    locator=f"slide {number}, {shape.name}",
                    evidence=evidence,
                )
            )
    return report


def apply(
    path: Path,
    fix_ids: list[str],
    output: Path,
    language: str = "en-US",
    title: str | None = None,
) -> dict:
    """Apply the named fixes and save to ``output``.

    Args:
        path: The original deck. Never overwritten.
        fix_ids: Fix identifiers: ``set-title`` and ``mark-header-row:<slide>:<shape>``.
        output: Where the remediated copy is written.
        language: Unused for decks; kept for a uniform signature.
        title: The title a person named for ``set-title``; without one, it is taken from
            the first slide title and the note says so.

    Returns:
        A dict with ``applied`` and ``failed``.
    """
    del language
    prs = Presentation(str(path))
    applied: dict[str, str] = {}
    failed: dict[str, str] = {}
    for fix_id in fix_ids:
        name, _, args = fix_id.partition(":")
        if name == "mark-header-row":
            try:
                slide_no, shape_index = (int(x) for x in args.split(":"))
                shape = list(prs.slides)[slide_no - 1].shapes[shape_index]
                if not shape.has_table:
                    raise ValueError("not a table")
                shape.table.first_row = True
                applied[fix_id] = f"slide {slide_no}: header row marked on {shape.name}"
            except (ValueError, IndexError) as exc:
                failed[fix_id] = f"{type(exc).__name__}: {exc}"
        elif fix_id == "set-title":
            if title:
                prs.core_properties.title = title
                applied[fix_id] = f"title set to {title!r}, as named"
                continue
            slide_titles = [
                s.shapes.title.text.strip()
                for s in prs.slides
                if s.shapes.title is not None and s.shapes.title.text.strip()
            ]
            if slide_titles:
                first, source = slide_titles[0], "the first slide title"
            else:
                first, source = path.stem.replace("_", " ").replace("-", " "), "the file name"
            prs.core_properties.title = first
            applied[fix_id] = (
                f"title set to {first!r} from {source}; check it names the whole presentation"
            )
        else:
            failed[fix_id] = "unknown fix"
    prs.save(str(output))
    return {"applied": applied, "failed": failed}
