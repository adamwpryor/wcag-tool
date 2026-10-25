"""The picture a finding points at, so alt text is drafted from looking, not guessing.

A scan says where a picture is (``Picture 1``, ``slide 4, Picture 2``, ``page 3``).
A model asked to draft alt text from that alone guesses from the file name and the
words nearby. ``extract`` returns the picture itself, as an image, using the finding's
own locator. It reads the file and writes nothing.

Drawings made of vector shapes, such as a logo drawn in a PDF, are not pictures in the
file's structure and cannot be returned; the person looks at those in the file.
"""

from __future__ import annotations

import re
from dataclasses import dataclass
from io import BytesIO
from pathlib import Path

from docx import Document
from PIL import Image, UnidentifiedImageError
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.shapes.base import BaseShape
from pptx.shapes.group import GroupShape
from pptx.shapes.picture import Picture as PptxPicture
from pypdf import PdfReader

from wcag_tool import limits
from wcag_tool.pdf_checks import drawn_images

# The longest side a returned picture keeps. Enough to judge what it shows; small
# enough that a pile of pictures does not swamp the model's context.
MAX_SIDE = 1568
# Above this a PNG is re-encoded as JPEG: textured artwork compresses badly as PNG, and
# model APIs commonly refuse images over about 5 MB.
_MAX_PNG_BYTES = 3_500_000

_SLIDE = re.compile(r"^slide (\d+), (.+)$", re.IGNORECASE)
_PAGES = re.compile(r"^pages? (\d+)(?:-(\d+))?$", re.IGNORECASE)


class AssetError(ValueError):
    """The locator names no picture this tool can return."""


@dataclass
class Picture:
    """One picture from a file.

    Attributes:
        label: Where it is, as a person would say it.
        data: The picture, scaled to ``MAX_SIDE`` at most.
        format: ``png``, or ``jpeg`` for artwork too heavy as PNG.
    """

    label: str
    data: bytes
    format: str = "png"


def _encode(data: bytes, where: str) -> tuple[bytes, str]:
    """Re-encode any picture Pillow can read, no larger than ``MAX_SIDE``."""
    try:
        with Image.open(BytesIO(data)) as image:
            image.load()
            picture = image.convert("RGBA") if image.mode in ("P", "LA", "PA") else image
            if picture.mode not in ("RGB", "RGBA", "L"):
                picture = picture.convert("RGB")
            picture.thumbnail((MAX_SIDE, MAX_SIDE))
            out = BytesIO()
            picture.save(out, format="PNG")
            if out.tell() <= _MAX_PNG_BYTES:
                return out.getvalue(), "png"
            out = BytesIO()
            picture.convert("RGB").save(out, format="JPEG", quality=85)
            return out.getvalue(), "jpeg"
    except (UnidentifiedImageError, OSError) as error:
        raise AssetError(
            f"The picture at {where} is in a format this tool cannot show (often a Windows "
            "metafile); open the file to look at it."
        ) from error


def _docx(path: Path, locator: str) -> list[Picture]:
    doc = Document(str(path))
    for doc_pr in doc.element.body.xpath(".//wp:docPr"):
        name = doc_pr.get("name") or f"image {doc_pr.get('id')}"
        if name != locator:
            continue
        rids = doc_pr.getparent().xpath(".//a:blip/@r:embed")
        if not rids:
            raise AssetError(f"{locator} has no embedded picture to show (it may be linked).")
        blob = doc.part.related_parts[rids[0]].blob
        return [Picture(locator, *_encode(blob, locator))]
    raise AssetError(f"No picture named {locator!r} in {path.name}.")


def _shapes(shapes: list[BaseShape]) -> list[BaseShape]:
    """Every shape on a slide, with the shapes inside groups."""
    out: list[BaseShape] = []
    for shape in shapes:
        out.append(shape)
        if isinstance(shape, GroupShape):
            out += _shapes(list(shape.shapes))
    return out


def _pptx(path: Path, locator: str) -> list[Picture]:
    match = _SLIDE.match(locator)
    if not match:
        raise AssetError(f"Expected a locator like 'slide 4, Picture 2', not {locator!r}.")
    number, name = int(match.group(1)), match.group(2)
    slides = Presentation(str(path)).slides
    if not 1 <= number <= len(slides):
        raise AssetError(f"{path.name} has no slide {number}.")
    for shape in _shapes(list(slides[number - 1].shapes)):
        if shape.name == name and shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
            assert isinstance(shape, PptxPicture)
            return [Picture(locator, *_encode(shape.image.blob, locator))]
    raise AssetError(f"No picture named {name!r} on slide {number} of {path.name}.")


def _pdf(path: Path, locator: str) -> list[Picture]:
    match = _PAGES.match(locator)
    if not match:
        raise AssetError(f"Expected a locator like 'page 3', not {locator!r}.")
    first = int(match.group(1))
    last = int(match.group(2) or first)
    reader = PdfReader(str(path))
    limits.check_pdf(reader, path.name)
    if not 1 <= first <= last <= len(reader.pages):
        raise AssetError(f"{path.name} has pages 1 to {len(reader.pages)}.")
    out: list[Picture] = []
    for number in range(first, last + 1):
        page = reader.pages[number - 1]
        wanted = [n for n in drawn_images(reader, page) if n is not None]
        # page.images lists every picture the page's resources name, drawn or not; keep the
        # ones this page draws, in drawing order.
        by_number = {
            image.indirect_reference.idnum: image
            for image in page.images
            if image.indirect_reference is not None
        }
        drawn = [by_number[n] for n in wanted if n in by_number]
        for index, image in enumerate(drawn, start=1):
            label = f"page {number}, picture {index} of {len(drawn)}"
            out.append(Picture(label, *_encode(image.data, label)))
    if not out:
        raise AssetError(
            f"No picture is drawn on {locator} of {path.name}. A drawing made of shapes (a "
            "logo, a diagram) is not a picture in the file; look at it in the file."
        )
    return out


def extract(path: Path, locator: str) -> list[Picture]:
    """The picture or pictures a finding's locator points at.

    Args:
        path: The course file. Read, never written.
        locator: A finding's ``locator``: ``Picture 1`` (Word), ``slide 4, Picture 2``
            (PowerPoint), ``page 3`` or ``pages 1-2`` (PDF).

    Returns:
        One picture, or for a PDF page every picture that page draws, in order.

    Raises:
        AssetError: If the locator names no picture this tool can return.
        limits.InputLimitError: If the file is over a size, package or page limit.
    """
    suffix = path.suffix.lower()
    if suffix in (".docx", ".pptx", ".pdf"):
        limits.preflight(path)
    if suffix == ".docx":
        return _docx(path, locator)
    if suffix == ".pptx":
        return _pptx(path, locator)
    if suffix == ".pdf":
        return _pdf(path, locator)
    raise AssetError(f"Unsupported file type {suffix!r}: expected .docx, .pptx or .pdf.")
