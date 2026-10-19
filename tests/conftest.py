import base64
from io import BytesIO
from pathlib import Path

import pymupdf
import pytest
from docx import Document
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from pptx import Presentation
from pptx.util import Inches

# A 1x1 transparent PNG. Generated fixtures only; no real course file is ever committed.
PNG_1X1 = base64.b64decode(
    "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAYAAAAfFcSJAAAADUlEQVR42mNkYPhfDwAChwGA60e6kgAAAABJRU5ErkJggg=="
)


@pytest.fixture
def bad_docx(tmp_path: Path) -> Path:
    """A Word file with one of every defect the checks look for."""
    doc = Document()
    doc.core_properties.title = ""
    for lang in doc.styles.element.xpath("./w:docDefaults/w:rPrDefault/w:rPr/w:lang"):
        lang.getparent().remove(lang)

    doc.add_paragraph("A course syllabus with no real structure.")
    fake = doc.add_paragraph()
    fake.add_run("Course Policies").bold = True
    doc.add_paragraph("Body text follows the bold line that is really a heading.")
    doc.add_picture(BytesIO(PNG_1X1))
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Week"
    table.cell(0, 1).text = "Topic"
    doc.add_heading("A real heading", level=1)
    doc.add_heading("Skips a level", level=3)

    link_paragraph = doc.add_paragraph("Read the policy: ")
    r_id = doc.part.relate_to("https://example.edu/policy", RT.HYPERLINK, is_external=True)
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), r_id)
    run = OxmlElement("w:r")
    text = OxmlElement("w:t")
    text.text = "click here"
    run.append(text)
    hyperlink.append(run)
    link_paragraph._p.append(hyperlink)

    path = tmp_path / "bad_syllabus.docx"
    doc.save(str(path))
    return path


@pytest.fixture
def bad_pptx(tmp_path: Path) -> Path:
    """A deck with an untitled slide and a picture without alt text."""
    prs = Presentation()
    prs.core_properties.title = ""
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide.shapes.add_picture(BytesIO(PNG_1X1), Inches(1), Inches(1), Inches(1), Inches(1))
    path = tmp_path / "bad_deck.pptx"
    prs.save(str(path))
    return path


@pytest.fixture
def bad_pdf(tmp_path: Path) -> Path:
    """An untagged PDF with no title and no language."""
    doc = pymupdf.open()
    page = doc.new_page()
    page.insert_text((72, 72), "Week 1 reading list")
    path = tmp_path / "bad_reading.pdf"
    doc.save(str(path))
    doc.close()
    return path
