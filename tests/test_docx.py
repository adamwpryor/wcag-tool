from pathlib import Path

import pytest
from docx import Document

from wcag_tool import engine
from wcag_tool.findings import Route


def _checks(report) -> set[str]:
    return {f.check for f in report.findings}


def test_scan_finds_every_defect(bad_docx: Path) -> None:
    report = engine.scan(bad_docx)
    assert _checks(report) >= {
        "docx.missing-title",
        "docx.missing-language",
        "docx.image-missing-alt",
        "docx.fake-heading",
        "docx.table-no-header",
        "docx.heading-skip",
        "docx.empty-heading",
        "docx.link-text",
    }
    assert report.evaluated and report.not_evaluated


def test_only_title_and_language_are_route_one(bad_docx: Path) -> None:
    report = engine.scan(bad_docx)
    route_one = {f.check for f in report.findings if f.route == Route.AUTO_FIX}
    assert route_one == {"docx.missing-title", "docx.missing-language"}
    for finding in report.findings:
        if finding.route == Route.CONTENT_OWNER:
            assert finding.proposal, finding.check


def test_fix_applies_route_one_only_by_default(bad_docx: Path) -> None:
    record = engine.fix(bad_docx, language="en-US")
    assert set(record["applied"]) == {"set-title", "set-language"}
    assert not record["failed"] and not record["held"] and not record["unknown_fix_ids"]
    after = {f["check"] for f in record["after"]["findings"]}
    assert "docx.missing-title" not in after
    assert "docx.missing-language" not in after
    assert "docx.table-no-header" in after
    # needs_person is everything still open, not only the findings with a fix id.
    assert {n["check"] for n in record["needs_person"]} == after
    assert {n["fix_id"] for n in record["needs_person"] if n["fix_id"]} == {
        "mark-header-row:0",
        "apply-heading:1:2",
    }
    assert Path(record["output"]).exists()
    assert Path(record["output"]) != bad_docx
    assert Path(record["output"]).with_suffix(".remediation.json").exists()


def test_fix_holds_language_until_a_person_names_it(bad_docx: Path) -> None:
    record = engine.fix(bad_docx)
    assert set(record["applied"]) == {"set-title"}
    assert record["language"] is None
    assert [h["fix_id"] for h in record["held"]] == ["set-language"]
    after = {f["check"] for f in record["after"]["findings"]}
    assert "docx.missing-language" in after


def test_fix_never_overwrites_the_original(bad_docx: Path) -> None:
    with pytest.raises(ValueError, match="never overwritten"):
        engine.fix(bad_docx, output=bad_docx)
    assert engine.scan(bad_docx).findings


def test_fix_applies_named_route_two_fix(bad_docx: Path) -> None:
    before = engine.scan(bad_docx)
    header_fix = next(f.fix_id for f in before.findings if f.check == "docx.table-no-header")
    heading_fix = next(f.fix_id for f in before.findings if f.check == "docx.fake-heading")
    record = engine.fix(bad_docx, apply=[header_fix, heading_fix, "bogus-fix"])
    assert header_fix in record["applied"] and heading_fix in record["applied"]
    assert record["unknown_fix_ids"] == ["bogus-fix"]
    after = {f["check"] for f in record["after"]["findings"]}
    assert "docx.table-no-header" not in after
    assert "docx.fake-heading" not in after
    fixed = Document(record["output"])
    assert fixed.core_properties.title == "Grading"


def test_confirmed_heading_level_may_differ_from_the_proposal(bad_docx: Path) -> None:
    record = engine.fix(bad_docx, apply=["apply-heading:1:1"])
    assert "apply-heading:1:1" in record["applied"]
    assert record["unknown_fix_ids"] == []
    fixed = Document(record["output"])
    assert fixed.paragraphs[1].style.name == "Heading 1"
